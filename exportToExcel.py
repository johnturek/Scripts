#!/usr/bin/env python3
"""
Export Cisco Catalyst Center data to Excel format with multiple worksheets.

This script creates a comprehensive Excel workbook with separate sheets for
switches, sites, and IP pools with formatted columns and better readability.
"""

import sys
import logging
import argparse
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any
from catalystcentersdk import CatalystCenter
from catalystcentersdk.exceptions import ApiError

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# Check if openpyxl is available
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    logger.warning("openpyxl not installed. Excel export will not be available.")
    logger.warning("Install with: pip install openpyxl")


def load_settings():
    """Load settings from settings.py file."""
    try:
        from settings import CATALYST_USERNAME, CATALYST_PASSWORD, CATALYST_BASE_URL
        return CATALYST_USERNAME, CATALYST_PASSWORD, CATALYST_BASE_URL
    except ImportError:
        logger.error("Settings file not found. Please create settings.py from settings.py.example")
        sys.exit(1)


def connect_to_catalyst(username: str, password: str, base_url: str) -> CatalystCenter:
    """Establish connection to Catalyst Center."""
    try:
        logger.info(f"Connecting to Catalyst Center at {base_url}...")
        catalyst = CatalystCenter(
            base_url=base_url,
            username=username,
            password=password,
            verify=True
        )
        logger.info("Successfully connected to Catalyst Center")
        return catalyst
    except Exception as e:
        logger.error(f"Failed to connect to Catalyst Center: {e}")
        sys.exit(1)


def get_switches_with_sites(catalyst: CatalystCenter) -> List[Dict[str, Any]]:
    """Retrieve all switches with site information."""
    try:
        logger.info("Retrieving switches with site information...")
        devices = catalyst.device.get_all_network_devices().response
        switches = [d for d in devices if 'Switch' in d.get('family', '')]
        
        site_cache = {}
        enriched_switches = []
        
        for idx, switch in enumerate(switches, 1):
            if idx % 10 == 0:
                logger.info(f"Processing switch {idx}/{len(switches)}...")
            
            site_id = switch.get('siteId')
            site_name = ''
            
            if site_id and site_id not in site_cache:
                try:
                    site_info = catalyst.site.get_site(site_id=site_id).response
                    site_cache[site_id] = site_info.get('site', {}).get('name', '')
                except:
                    site_cache[site_id] = ''
            
            if site_id:
                site_name = site_cache.get(site_id, '')
            
            enriched_switches.append({
                'hostname': switch.get('hostname', ''),
                'managementIpAddress': switch.get('managementIpAddress', ''),
                'macAddress': switch.get('macAddress', ''),
                'serialNumber': switch.get('serialNumber', ''),
                'platformId': switch.get('platformId', ''),
                'softwareVersion': switch.get('softwareVersion', ''),
                'role': switch.get('role', ''),
                'site_name': site_name,
                'reachabilityStatus': switch.get('reachabilityStatus', ''),
                'upTime': switch.get('upTime', ''),
            })
        
        logger.info(f"Retrieved {len(enriched_switches)} switches")
        return enriched_switches
    except ApiError as e:
        logger.error(f"API error: {e}")
        return []


def get_sites(catalyst: CatalystCenter) -> List[Dict[str, Any]]:
    """Retrieve all sites."""
    try:
        logger.info("Retrieving all sites...")
        sites_response = catalyst.site.get_all_sites().response
        
        sites_data = []
        for site in sites_response:
            site_info = site.get('site', {})
            sites_data.append({
                'name': site_info.get('name', ''),
                'siteId': site.get('id', ''),
                'parentId': site.get('parentId', ''),
                'latitude': site_info.get('latitude', ''),
                'longitude': site_info.get('longitude', ''),
                'locationType': site_info.get('locationType', ''),
            })
        
        logger.info(f"Retrieved {len(sites_data)} sites")
        return sites_data
    except ApiError as e:
        logger.error(f"API error: {e}")
        return []


def get_ip_pools(catalyst: CatalystCenter) -> List[Dict[str, Any]]:
    """Retrieve all IP pools."""
    try:
        logger.info("Retrieving all IP pools...")
        pools_response = catalyst.network_settings.get_all_ip_pools().response
        
        pools_data = []
        for pool in pools_response:
            pools_data.append({
                'ipPoolName': pool.get('ipPoolName', ''),
                'ipPoolCidr': pool.get('ipPoolCidr', ''),
                'gateway': pool.get('gateway', ''),
                'dhcpServerIps': ', '.join(pool.get('dhcpServerIps', [])),
                'dnsServerIps': ', '.join(pool.get('dnsServerIps', [])),
                'IpAddressSpace': pool.get('IpAddressSpace', ''),
            })
        
        logger.info(f"Retrieved {len(pools_data)} IP pools")
        return pools_data
    except ApiError as e:
        logger.error(f"API error: {e}")
        return []


def format_worksheet(ws, data: List[Dict], title: str):
    """Apply formatting to worksheet."""
    if not data:
        return
    
    # Style for headers
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Write headers
    headers = list(data[0].keys())
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    # Write data
    for row_num, row_data in enumerate(data, 2):
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = row_data.get(header, '')
    
    # Auto-adjust column widths
    for col_num in range(1, len(headers) + 1):
        column_letter = get_column_letter(col_num)
        max_length = len(headers[col_num - 1])
        
        for row in ws.iter_rows(min_row=2, max_row=len(data) + 1, min_col=col_num, max_col=col_num):
            for cell in row:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
        
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Freeze the header row
    ws.freeze_panes = 'A2'


def export_to_excel(switches: List[Dict], sites: List[Dict], ip_pools: List[Dict], filename: str):
    """Export data to Excel workbook with multiple sheets."""
    if not EXCEL_AVAILABLE:
        logger.error("Excel export requires openpyxl. Install with: pip install openpyxl")
        sys.exit(1)
    
    try:
        logger.info(f"Creating Excel workbook: {filename}")
        wb = Workbook()
        
        # Remove default sheet
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        
        # Add switches sheet
        if switches:
            ws_switches = wb.create_sheet("Switches")
            format_worksheet(ws_switches, switches, "Network Switches")
            logger.info(f"Added {len(switches)} switches to worksheet")
        
        # Add sites sheet
        if sites:
            ws_sites = wb.create_sheet("Sites")
            format_worksheet(ws_sites, sites, "Sites")
            logger.info(f"Added {len(sites)} sites to worksheet")
        
        # Add IP pools sheet
        if ip_pools:
            ws_pools = wb.create_sheet("IP Pools")
            format_worksheet(ws_pools, ip_pools, "IP Pools")
            logger.info(f"Added {len(ip_pools)} IP pools to worksheet")
        
        # Add summary sheet
        ws_summary = wb.create_sheet("Summary", 0)
        ws_summary['A1'] = "Catalyst Center Export Summary"
        ws_summary['A1'].font = Font(bold=True, size=14)
        ws_summary['A3'] = "Export Date:"
        ws_summary['B3'] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws_summary['A4'] = "Total Switches:"
        ws_summary['B4'] = len(switches)
        ws_summary['A5'] = "Total Sites:"
        ws_summary['B5'] = len(sites)
        ws_summary['A6'] = "Total IP Pools:"
        ws_summary['B6'] = len(ip_pools)
        
        # Save workbook
        wb.save(filename)
        logger.info(f"Successfully exported data to {filename}")
        
    except Exception as e:
        logger.error(f"Failed to create Excel file: {e}")
        sys.exit(1)


def main():
    """Main execution function."""
    parser = argparse.ArgumentParser(
        description='Export Cisco Catalyst Center data to Excel format'
    )
    parser.add_argument(
        '-o', '--output',
        default='catalyst_export.xlsx',
        help='Output Excel file name (default: catalyst_export.xlsx)'
    )
    parser.add_argument(
        '-v', '--verbose',
        action='store_true',
        help='Enable verbose logging'
    )
    
    args = parser.parse_args()
    
    if args.verbose:
        logger.setLevel(logging.DEBUG)
    
    if not EXCEL_AVAILABLE:
        logger.error("This script requires openpyxl. Install with: pip install openpyxl")
        sys.exit(1)
    
    # Load configuration
    username, password, base_url = load_settings()
    
    # Connect to Catalyst Center
    catalyst = connect_to_catalyst(username, password, base_url)
    
    # Retrieve data
    switches = get_switches_with_sites(catalyst)
    sites = get_sites(catalyst)
    ip_pools = get_ip_pools(catalyst)
    
    if not switches and not sites and not ip_pools:
        logger.warning("No data found to export")
        sys.exit(0)
    
    # Export to Excel
    export_to_excel(switches, sites, ip_pools, args.output)
    
    logger.info("Export completed successfully!")


if __name__ == "__main__":
    main()
